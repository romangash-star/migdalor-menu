# -*- coding: utf-8 -*-
"""ממיר ייצוא Excel של לוח "קול קורא לתפריט רשויות" לפריטים במבנה של האתר.

שימוש:
    python import_kol_kore.py --excel "קול קורא לתפריט רשויות.xlsx" --out new_items.json

הייצוא מגיע עם שורת כותרת בשורה 3, ועם שורות-כותרת נוספות שמפרידות בין
הקבוצות בלוח; הן מזוהות ומדולגות.
"""

import argparse
import json
import io
import re

CATEGORY = {
    'תכניות חינוכיות וקהילתיות': 'תכניות חינוכיות וקהילתיות',
    'חזון ותפיסת יעוד':          'חזון ויעוד',
    'כח אדם':                    'כח אדם',
    'תרבות ארגונים':             'תרבות ארגונית',
    'שותפויות ומשאבים':          'שותפויות ומשאבים',
}
TOPIC = {'יחסים בין בקבוצות': 'יחסים בין קבוצות'}

BULLET = re.compile('^[\\s•·●▪*\\-–—]+')
CONTROL = re.compile('[\\x00-\\x1f]+')


def clean_name(s):
    """חלק מהשמות בלוח מתחילים בתו תבליט ובטאב; באתר הם צריכים להופיע נקיים."""
    s = BULLET.sub('', s or '')
    s = CONTROL.sub(' ', s)
    return re.sub(r'\s{2,}', ' ', s).strip()


def phone(v):
    """972501234567 -> 050-123-4567, כדי שגם החיוג וגם הוואטסאפ יעבדו."""
    d = re.sub(r'\D', '', v or '')
    if d.startswith('972'):
        d = '0' + d[3:]
    if len(d) == 10 and d.startswith('0'):
        return '%s-%s-%s' % (d[:3], d[3:6], d[6:])
    if len(d) == 9 and d.startswith('0'):
        return '%s-%s-%s' % (d[:2], d[2:5], d[5:])
    return (v or '').strip()


def doc_urls(v):
    """תא אחד יכול להחזיק כמה קישורים ברצף, בלי מפריד."""
    parts = re.split('(?=https?://)', v or '')
    return [u.strip().rstrip(',;') for u in parts if u.strip().startswith('http')]


def read_rows(excel):
    import openpyxl
    ws = openpyxl.load_workbook(excel, data_only=True).worksheets[0]
    hdr = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(4, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not any(v not in (None, '') for v in vals):
            continue
        row = {hdr[i]: (str(vals[i]).strip() if vals[i] is not None else '')
               for i in range(len(hdr))}
        if row.get('Name') == 'Name':          # שורת כותרת שחוזרת בין הקבוצות
            continue
        rows.append(row)
    return rows


def convert(rows):
    items, skipped = [], []
    for r in rows:
        name = clean_name(r.get('Name'))
        cat = CATEGORY.get((r.get('תחום בתפריט') or '').strip())
        if not name or not cat:
            skipped.append({'name': name, 'why': 'תחום לא מזוהה: ' + (r.get('תחום בתפריט') or '(ריק)')})
            continue
        topic = (r.get('נושא מרכזי של הפריט המוצע') or '').strip()
        items.append({
            'name': name,
            'category': cat,
            'dept': (r.get('מחלקה רלבנטית') or '').strip(),
            'org': (r.get('שם הארגון/הספק המבצע') or '').strip(),
            'contactName': (r.get('שם איש הקשר') or '').strip(),
            'phone': phone(r.get('טלפון ליצירת קשר')),
            'email': (r.get('מייל ליצירת קשר') or '').strip(),
            'audience': ', '.join(a.strip() for a in re.split(r'\s*,\s*', r.get('קהל יעד') or '') if a.strip()),
            'topic': TOPIC.get(topic, topic),
            'gafan': (r.get('האם הפריט קיים במאגר גפ"ן') or '').strip(),
            'type': (r.get('סוג פריט') or 'ספק').strip(),
            'description': (r.get('תיאור ההצעה') or '').strip(),
            'notes': (r.get('דגשים/הערות') or '').strip(),
            'website': (r.get('לינק לאתר הספק/הפעילות') or '').strip(),
            'files': [],
            '_docs': doc_urls(r.get('מסמכים רלוונטיים')),
        })
    return items, skipped


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--excel', required=True)
    ap.add_argument('--out', default='new_items.json')
    args = ap.parse_args()

    items, skipped = convert(read_rows(args.excel))
    io.open(args.out, 'w', encoding='utf-8').write(json.dumps(items, ensure_ascii=False, indent=1))
    print('הומרו %d פריטים אל %s' % (len(items), args.out))
    print('עם מסמכים מצורפים: %d' % sum(1 for i in items if i['_docs']))
    for s in skipped:
        print('   דולג:', s['name'][:40], '|', s['why'])


if __name__ == '__main__':
    main()
